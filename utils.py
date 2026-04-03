from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Dict, List

from pptx import Presentation


def load_presentation(pptx_path: str) -> Presentation:
    return Presentation(pptx_path)


def normalize_text(text: str) -> str:
    """
    Normalize extracted text so minor formatting changes don't create noise.
    """
    if text is None:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    # Collapse excessive blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def shape_has_text_frame(shape: Any) -> bool:
    return bool(getattr(shape, "has_text_frame", False))


def extract_text_from_table(shape: Any) -> str:
    table = shape.table
    rows_text: List[str] = []

    for r in range(len(table.rows)):
        cols_text: List[str] = []
        for c in range(len(table.columns)):
            cell_text = table.cell(r, c).text or ""
            cell_text = normalize_text(cell_text)
            cols_text.append(cell_text)
        rows_text.append(" | ".join(cols_text).strip())

    return normalize_text("\n".join(rows_text))


def extract_text_from_text_frame(shape: Any) -> str:
    tf = shape.text_frame
    parts: List[str] = []
    for paragraph in tf.paragraphs:
        p = normalize_text(paragraph.text)
        if p:
            parts.append(p)
    return "\n".join(parts).strip()


def extract_text_from_group_like(shape: Any) -> str:
    """
    Some shapes (e.g., group shapes) may contain nested shapes via `.shapes`.
    """
    if not hasattr(shape, "shapes"):
        return ""

    parts: List[str] = []
    try:
        for child in shape.shapes:
            child_text = extract_text_from_shape(child)
            if child_text:
                parts.append(child_text)
    except Exception:
        # Best effort only; don't fail whole extraction.
        return ""

    return normalize_text("\n".join(parts))


def extract_text_from_shape(shape: Any) -> str:
    """
    Extract visible text from a single shape.
    For non-text shapes, returns "".
    """
    # Table
    if getattr(shape, "has_table", False):
        try:
            return extract_text_from_table(shape)
        except Exception:
            return ""

    # Text frame
    if shape_has_text_frame(shape):
        try:
            return extract_text_from_text_frame(shape)
        except Exception:
            return ""

    # Group-like container
    group_text = extract_text_from_group_like(shape)
    if group_text:
        return group_text

    return ""


def extract_slide_model(slide: Any) -> Dict[str, Any]:
    """
    Build a comparable slide model.
    MVP: concatenate all extracted shape texts in on-slide order.
    """
    shape_texts: List[str] = []

    for shape in slide.shapes:
        txt = extract_text_from_shape(shape)
        if txt:
            shape_texts.append(txt)

    slide_text = normalize_text("\n\n".join(shape_texts))

    return {
        "slide_number": None,  # caller fills
        "text": slide_text,
        "shape_texts": shape_texts,  # helpful for debugging
    }


def extract_presentation_model(prs: Presentation) -> Dict[str, Any]:
    slides: List[Dict[str, Any]] = []
    for idx, slide in enumerate(prs.slides):
        m = extract_slide_model(slide)
        m["slide_number"] = idx
        slides.append(m)

    return {
        "slide_count": len(slides),
        "slides": slides,
    }


def write_diff_report(report: Dict[str, Any], out_path: str) -> None:
    """
    Save the diff report JSON to `out_path`.
    """
    path = Path(out_path)
    # Avoid mkdir(".") when out_path is just a filename in CWD.
    if path.parent not in (Path("."), Path("")):
        path.parent.mkdir(parents=True, exist_ok=True)

    path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

